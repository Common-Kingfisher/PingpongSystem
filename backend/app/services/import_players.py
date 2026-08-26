"""Excel / CSV 选手批量导入（Demo，仅 .xlsx / .csv）。

- 第一行为表头，支持简单别名匹配（去空格、英文忽略大小写）；
- 逐行读取：空行忽略、姓名为空跳过、姓名/单位 trim；
- 种子序号：>=1、不重复、不超小组数；违规时选手仍导入但种子不生效；
- 不做姓名去重（模型无唯一约束，真实赛事允许同名）；
- 文件级错误（无姓名列/无法解码/格式不支持/赛事锁定）整体失败。
"""

import csv
import io
import sqlite3

from .. import repository as repo
from ..models import TournamentStage

NAME_ALIASES = {"姓名", "选手姓名", "名字", "name", "player_name"}
COLLEGE_ALIASES = {"学院", "学院/单位", "单位", "学校", "部门", "organization", "college"}
SEED_ALIASES = {"种子", "种子序号", "种子编号", "seed", "seed_no"}


class ImportFileError(Exception):
    def __init__(self, message: str, code: int = 400):
        super().__init__(message)
        self.code = code


def _norm(header: str) -> str:
    return header.strip().lstrip("\ufeff").lower()


def _find_column(headers: list[str], aliases: set[str]) -> int | None:
    alias_set = {_norm(a) for a in aliases}
    for i, h in enumerate(headers):
        if _norm(h) in alias_set:
            return i
    return None


def _cell(row: list[str], i: int | None) -> str:
    return row[i].strip() if i is not None and i < len(row) else ""


def _parse_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")  # 兼容 UTF-8 与 UTF-8 BOM
    except UnicodeDecodeError:
        raise ImportFileError("CSV 文件编码无法识别，请使用 UTF-8 编码重新保存。")
    rows = []
    for row in csv.reader(io.StringIO(text)):
        cells = [c.strip() for c in row]
        if any(cells):
            rows.append(cells)
    return rows


def _parse_xlsx(content: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportFileError("服务端缺少 openpyxl 依赖，无法解析 .xlsx 文件", 500)
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]  # 只读第一个工作表
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                rows.append(cells)
        wb.close()
    except Exception:
        raise ImportFileError("无法打开 Excel 文件，请确认其为有效的 .xlsx 文件")
    return rows


def import_players_file(
    conn: sqlite3.Connection, tournament_id: int, content: bytes, filename: str
) -> dict:
    tournament = repo.get_tournament(conn, tournament_id)
    if tournament is None:
        raise ImportFileError("赛事不存在", 404)
    if tournament["stage"] != TournamentStage.REGISTRATION.value:
        raise ImportFileError("赛事已进入比赛阶段，选手名单已锁定", 409)

    lower = filename.lower()
    if lower.endswith(".csv"):
        rows = _parse_csv(content)
    elif lower.endswith(".xlsx"):
        rows = _parse_xlsx(content)
    else:
        raise ImportFileError("不支持的文件格式，请使用 .xlsx 或 .csv")

    if not rows:
        raise ImportFileError("文件为空")

    headers = rows[0]
    name_col = _find_column(headers, NAME_ALIASES)
    if name_col is None:
        raise ImportFileError(
            "未找到“姓名”列。请使用以下任一表头：姓名 / 选手姓名 / 名字 / name"
        )
    college_col = _find_column(headers, COLLEGE_ALIASES)
    seed_col = _find_column(headers, SEED_ALIASES)

    existing = repo.list_players(conn, tournament_id)
    used_seeds = {p["seed_no"] for p in existing if p["seed_no"] is not None}
    max_seed = tournament["group_count"]

    errors: list[dict] = []
    imported = 0
    skipped = 0
    total_rows = 0

    for row_no, row in enumerate(rows[1:], start=2):
        total_rows += 1
        name = _cell(row, name_col)
        if not name:
            skipped += 1
            errors.append({"row": row_no, "message": "姓名为空"})
            continue
        college = _cell(row, college_col) or None

        seed_no: int | None = None
        seed_raw = _cell(row, seed_col)
        if seed_raw:
            try:
                seed_no = int(seed_raw)
            except ValueError:
                seed_no = None
                errors.append({"row": row_no, "message": "种子序号格式错误，选手已导入但未设置种子"})
            if seed_no is not None:
                if seed_no < 1:
                    errors.append({"row": row_no, "message": "种子序号必须 >= 1，选手已导入但未设置种子"})
                    seed_no = None
                elif seed_no > max_seed:
                    errors.append(
                        {"row": row_no, "message": f"种子序号超出小组数（最多 {max_seed}），选手已导入但未设置种子"}
                    )
                    seed_no = None
                elif seed_no in used_seeds:
                    errors.append(
                        {"row": row_no, "message": f"{seed_no}号种子已被其他选手占用，选手已导入但未设置种子"}
                    )
                    seed_no = None

        player = repo.add_player(conn, tournament_id, name, college)
        if seed_no is not None:
            repo.set_player_seed(conn, player["id"], seed_no)
            used_seeds.add(seed_no)
        imported += 1

    conn.commit()
    return {
        "total_rows": total_rows,
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
    }
